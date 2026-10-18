"""pin-map.xlsx for the STM32G474RBT6 (LQFP64) on the DCDC LLC + 4SBB board (Rev D).
Pin names, types, AF0..AF15 and additional functions come from ST's pin database (STM32_open_pin_data:
mcu/STM32G474R(B-C-E)Tx.xml + mcu/IP/GPIO-STM32G47x_gpio_v1_0_Modes.xml), which mirrors datasheet DS12288
Table 13 'pin definitions' and Table 14/15 'alternate functions'. Design columns come from design.py (U1 pins)."""
import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, "/home/claude/dcdc")
import design as D

PINS = json.load(open(sys.argv[1]))
OUT = sys.argv[2]
U1 = next(p for p in D.P if p.ref == "U1")

# net -> who else is on it
others = {}
for p in D.P:
    for pin, n in p.pins.items():
        if p.ref != "U1":
            others.setdefault(n, []).append(f"{p.ref}.{pin}")

# ---- what the design uses on each pin: (function used, mode / AF, direction, notes)
USE = {
    "+3V3": ("VDD supply", "Power", "Supply", "100 nF per VDD pin (C30-C34) + 4.7 uF bulk (C35)"),
    "+3V3A": ("Analog supply / reference", "Power", "Supply", "VDDA and VREF+ from +3V3 via FB1; VREFBUF off, VREF+ used as input (1 uF + 100 nF)"),
    "GND": ("Ground", "Power", "Supply", "Secondary GND plane (In1)"),
    "OSC_IN": ("RCC_OSC_IN (HSE)", "RCC HSE crystal", "In", "Y1 16 MHz 9 pF + 12 pF C0G load caps -> PLL 170 MHz"),
    "OSC_OUT": ("RCC_OSC_OUT (HSE)", "RCC HSE crystal", "Out", "Y1 16 MHz"),
    "NRST": ("NRST", "Reset", "In/Out", "100 nF to GND (C40); SWD reset on J6.6 / J7.10"),
    "BOOT0": ("BOOT0", "Boot pin", "In", "10k to GND (R20) -> boot from flash; nSWBOOT0 option bit may override"),
    "VBUS_SNS": ("ADC1_IN1", "Analog", "In", "VBUS divider 100k/5.1k, 1 nF -> 68 V full scale"),
    "IOUT_SNS": ("ADC1_IN2", "Analog", "In", "INA181A2 on 10 mOhm (RSH1): 0.5 V/A, 1k + 10 nF filter"),
    "NTC_SNS": ("ADC2_IN3", "Analog", "In", "NTC 10k B3950 / 10k divider (Rev D: moved from PA4 to PA6)"),
    "VOUT_SNS": ("ADC1_IN14 + COMP6_INP", "Analog", "In", "VOUT divider 100k/5.1k; COMP6 = fast hardware OVP (threshold from DAC3)"),
    "TP_DAC": ("DAC1_OUT2", "Analog", "Out", "Test_pin_DAC on PROG PADS J6.3 (debug waveform out)"),
    "LLC_A": ("HRTIM1_CHC1", "AF13", "Out", "LLC half-bridge high side via CA-IS3642 (U12) -> UCC27211 (U10)"),
    "LLC_B": ("HRTIM1_CHC2", "AF13", "Out", "LLC half-bridge low side via U12 -> U10"),
    "SR1_PWM": ("HRTIM1_CHD1", "AF13", "Out", "Synchronous rectifier Q7 via UCC27524 (U11)"),
    "SR2_PWM": ("HRTIM1_CHD2", "AF13", "Out", "Synchronous rectifier Q8 via U11"),
    "PWM_E1": ("HRTIM1_CHE1", "AF3", "Out", "4SBB output leg high side Q3 via UCC27211 (U3)"),
    "PWM_E2": ("HRTIM1_CHE2", "AF3", "Out", "4SBB output leg low side Q4 via U3"),
    "PWM_A1": ("HRTIM1_CHA1", "AF13", "Out", "4SBB input leg high side Q1 via UCC27211 (U2)"),
    "PWM_A2": ("HRTIM1_CHA2", "AF13", "Out", "4SBB input leg low side Q2 via U2"),
    "CAN_RXD": ("FDCAN1_RX", "AF9", "In", "CA-IS3062VW (U9) RXD"),
    "CAN_TXD": ("FDCAN1_TX", "AF9", "Out", "CA-IS3062VW (U9) TXD"),
    "SWDIO": ("SYS_JTMS-SWDIO", "AF0 (reset default)", "In/Out", "J6.1 / J7.2"),
    "SWCLK": ("SYS_JTCK-SWCLK", "AF0 (reset default)", "In", "J6.2 / J7.4"),
    "SWO": ("SYS_JTDO-SWO", "AF0 (reset default)", "Out", "J7.6 trace output"),
    "VIN_OVP": ("HRTIM1_FLT2 (alt.: GPIO EXTI15)", "AF13", "In", "Rev D: primary input OVP, LMV331 + TL431 trip 50.2 V, through U12 VI1->VO1; "
                "high = over-voltage. Configure FLT2 active-high so the HRTIM shuts all outputs in hardware"),
    "LED_FLT": ("GPIO output", "Output push-pull", "Out", "Red FAULT LED D5 via 1k, active low"),
    "LED_RUN": ("GPIO output", "Output push-pull", "Out", "Green RUN LED D4 via 1k, active low"),
    "SR_EN": ("GPIO output", "Output push-pull", "Out", "UCC27524 ENA/ENB, high = SR drive enabled; 10k pull-downs keep it off at reset"),
    "TP_GPIO1": ("GPIO", "Output (debug)", "In/Out", "Test_pin_GPIO1 on J6.4 (timing / debug)"),
    "TP_GPIO2": ("GPIO", "Output (debug)", "In/Out", "Test_pin_GPIO2 on J6.5 (timing / debug)"),
}
for i in range(6):
    USE[f"ADDR{i}"] = ("GPIO input", "Input, no pull (10k external pull-up)", "In", f"Address bit {i}: DIP SW1 / header J5 (ON or jumper = 0)")
SPECIAL = {
    "IL_SNS@24": ("COMP4_INP", "Analog", "In", "Inductor current (INA181A1 on 4 mOhm, 1.65 V + 0.08 V/A): buck-mode negative OCP, COMP4 output inverted -> HRTIM fault"),
    "IL_SNS@25": ("ADC1_IN12 + COMP1_INP", "Analog", "In", "Inductor current sample for CC loop; COMP1 = positive OCP -> HRTIM fault (thresholds from DAC3)"),
}

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
F = lambda **k: Font(name="Arial", size=9, **k)
FILL = {"used": PatternFill("solid", fgColor="E2EFDA"), "power": PatternFill("solid", fgColor="FFF2CC"),
        "unused": PatternFill("solid", fgColor="EDEDED"), "reserved": PatternFill("solid", fgColor="DDEBF7"),
        "head": PatternFill("solid", fgColor="1F3864")}

wb = Workbook()
ws = wb.active
ws.title = "Pin map"
heads = ["Pin #", "Pin name", "Pin type", "Net (schematic)", "Status", "Function used", "Mode / AF", "Direction",
         "Connected to", "Design notes", "Additional functions (analog, wake-up, ...)"] + [f"AF{i}" for i in range(16)]
ws.append(heads)
for c in range(1, len(heads) + 1):
    cell = ws.cell(row=1, column=c)
    cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
    cell.fill = FILL["head"]
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    cell.border = BORDER

for d in PINS:
    pos = d["pos"]
    net = U1.pins.get(str(pos), "")
    if net in ("+3V3", "+3V3A", "GND"):
        status = "Power"
    elif net.startswith("SP_"):
        status = "Unused"
    else:
        status = "Used"
    if net.startswith("SP_"):
        port = net[3:]
        rref = next(p.ref for p in D.P if p.pins.get("1") == net)
        func, mode, dirn = "Not used", "Analog (reset state) or input", "-"
        note = f"10k pull-down to GND ({rref})"
        if port == "PA4":
            status = "Reserved"
            func, note = "Reserved: DAC1_OUT1", f"Kept free for DAC1_OUT1 (Rev D request); 10k pull-down {rref} is a valid DAC load (min 5k)"
        if port == "PB4":
            note += "; PB4 = NJTRST with internal pull-up at reset - harmless, JTAG unused (SWD only)"
        if port in ("PC13", "PC14", "PC15"):
            note += "; low-drive backup-domain pin"
    elif d["name"] == "VBAT":
        func, mode, dirn, note = "VBAT (backup domain)", "Power", "Supply", "Tied to +3V3 (no backup battery; RTC/backup registers lost at power-off)"
    elif d["name"] == "VREF+":
        func, mode, dirn, note = "VREF+ input", "Power", "Supply", "Tied to +3V3A (VREFBUF disabled); ADC/DAC full scale = 3.3 V"
    elif f"{net}@{pos}" in SPECIAL:
        func, mode, dirn, note = SPECIAL[f"{net}@{pos}"]
    else:
        func, mode, dirn, note = USE.get(net, ("?", "?", "?", ""))
    conn = ", ".join(sorted(others.get(net, []))) if status != "Power" else ""
    af = {v: [] for v in range(16)}
    for sig, n in d["af"].items():
        af[n].append(sig)
    addl = ", ".join(s for s in d["sigs"] if s not in d["af"] and s != "GPIO")
    row = [pos, d["name"], d["type"], net, status, func, mode, dirn, conn, note, addl] + [" / ".join(af[i]) for i in range(16)]
    ws.append(row)
    r = ws.max_row
    fill = FILL["power"] if status == "Power" else FILL["unused"] if status == "Unused" else FILL["reserved"] if status == "Reserved" else FILL["used"]
    for c in range(1, len(heads) + 1):
        cell = ws.cell(row=r, column=c)
        cell.font = F(bold=(c in (6, 7) and status in ("Used", "Reserved")))
        cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if c in (1, 3, 5, 8) else "left")
        cell.border = BORDER
        if c <= 10:
            cell.fill = fill

widths = [6, 16, 8, 13, 9, 24, 17, 9, 26, 46, 40] + [14] * 16
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 30
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{ws.max_row}"
last = ws.max_row

# ---- summary sheet (formulas)
s = wb.create_sheet("Summary")
s["A1"] = "STM32G474RBT6 (LQFP64) pin usage - DCDC LLC + 4SBB, Rev D"
s["A1"].font = Font(name="Arial", size=12, bold=True)
rows = [("Status", "Pins", None), ("Used", f'=COUNTIF(\'Pin map\'!$E$2:$E${last},"Used")', None),
        ("Reserved", f'=COUNTIF(\'Pin map\'!$E$2:$E${last},"Reserved")', None),
        ("Unused (10k pull-down)", f'=COUNTIF(\'Pin map\'!$E$2:$E${last},"Unused")', None),
        ("Power / ground", f'=COUNTIF(\'Pin map\'!$E$2:$E${last},"Power")', None),
        ("Total", "=SUM(B4:B7)", None),
        ("Check (must be 64)", f"=COUNTA('Pin map'!$A$2:$A${last})", None)]
for i, (a, b, _) in enumerate(rows):
    s.cell(row=3 + i, column=1, value=a).font = F(bold=(i in (0, 5)))
    s.cell(row=3 + i, column=2, value=b).font = F(bold=(i in (0, 5)))
s["A12"] = "Peripheral resources in use"
s["A12"].font = Font(name="Arial", size=10, bold=True)
per = [("Peripheral", "Signals (pin)"),
       ("HRTIM1 Timer A", "CHA1 PA8, CHA2 PA9 (AF13) - 4SBB input leg, 150 kHz"),
       ("HRTIM1 Timer E", "CHE1 PC8, CHE2 PC9 (AF3) - 4SBB output leg, 150 kHz"),
       ("HRTIM1 Timer C", "CHC1 PB12, CHC2 PB13 (AF13) - LLC half-bridge, ~205 kHz"),
       ("HRTIM1 Timer D", "CHD1 PB14, CHD2 PB15 (AF13) - synchronous rectifiers"),
       ("HRTIM1 faults", "FLT2 PA15 (AF13) = input OVP; internal COMP1/COMP4/COMP6 outputs as fault sources"),
       ("ADC1", "IN1 PA0 VBUS, IN2 PA1 IOUT, IN12 PB1 IL, IN14 PB11 VOUT"),
       ("ADC2", "IN3 PA6 NTC"),
       ("COMP1 / COMP4 / COMP6", "INP PB1 (+OCP), INP PB0 (-OCP), INP PB11 (VOUT OVP); INM from DAC3 channels"),
       ("DAC1", "OUT2 PA5 = Test_pin_DAC; OUT1 PA4 reserved"),
       ("DAC3 (internal)", "Comparator thresholds, no pins"),
       ("FDCAN1", "RX PA11, TX PA12 (AF9) -> isolated CAN"),
       ("SWD", "SWDIO PA13, SWCLK PA14, SWO PB3, NRST"),
       ("RCC", "HSE 16 MHz crystal PF0/PF1 -> PLL 170 MHz"),
       ("GPIO", "ADDR0-5 PC0-PC5 (in), LED_RUN PB9, LED_FLT PC10, SR_EN PC11, TP_GPIO1/2 PB6/PB7")]
for i, (a, b) in enumerate(per):
    s.cell(row=13 + i, column=1, value=a).font = F(bold=(i == 0))
    s.cell(row=13 + i, column=2, value=b).font = F(bold=(i == 0))
    for c in (1, 2):
        s.cell(row=13 + i, column=c).border = BORDER
        s.cell(row=13 + i, column=c).alignment = Alignment(wrap_text=True, vertical="top")
for r in range(3, 10):
    for c in (1, 2):
        s.cell(row=r, column=c).border = BORDER
s.column_dimensions["A"].width = 24
s.column_dimensions["B"].width = 90
s["A30"] = "Legend: green = used, blue = reserved, grey = unused (10k pull-down), yellow = power. Pin-map columns K-AA copied from the ST pin database."
s["A31"] = ("Source: STMicroelectronics STM32_open_pin_data (CubeMX DB) - mcu/STM32G474R(B-C-E)Tx.xml and "
            "mcu/IP/GPIO-STM32G47x_gpio_v1_0_Modes.xml; matches datasheet DS12288 pin definition and alternate-function tables.")
s["A32"] = "Design columns D-J generated from the Rev D netlist (design.py, U1). Edit the netlist, not this sheet, and regenerate."
for a in ("A30", "A31", "A32"):
    s[a].font = F(italic=True)
wb.save(OUT)
print("saved", OUT, "rows", last - 1)
