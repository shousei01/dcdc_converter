"""Rev E sourcing review: LCSC stock today (live, lcsc.com product pages, 2026-10-20 ~09:00 CST),
shortage risk, 2nd / 3rd sources. Output: 01_hardware/bom/BOM_sourcing_revE.xlsx"""
import csv, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
sys.path.insert(0, "sourcing")
from alts import ALTS

STOCK = {}
for r in csv.reader(open("sourcing/lcsc_stock_20261020.tsv"), delimiter="\t"):
    r = r + [""] * (5 - len(r)); STOCK[r[0]] = (r[1], r[2], r[3], r[4])
bom = list(csv.reader(open("release_E/jlc/DCDC_JLC_BOM.csv")))[1:]
bom.append(["STM32G474RBT6", "U1", "LQFP-64", "C529413"])     # hand-soldered, but the key shortage item

# line-specific findings (action, risk override, note)
FIX = {
 "C529413": ("CRITICAL", "Out of stock at LCSC. Buy STM32G474RBT3 (C1235414, same die, -40..125 C, 564 pcs) now; RET6 has only 12. Firmware unchanged (RB = 128 KB)."),
 "C503012": ("CRITICAL", "Only 38 pcs, 4 per board. Switch Q2/Q3/Q5/Q6 to NCEP068N10G (C503061, 7,577 pcs, 6.1 mOhm, Qg 60 nC): check 4SBB dead time / driver loss (~0.1 W per FET at 150 kHz, fine for UCC27211)."),
 "C1525": ("CRITICAL", "Out of stock. Move these 22 caps onto the C307331 line (same 100 nF 0402 X7R, 50 V, 2.3 M pcs) - one less BOM line."),
 "C25744": ("CRITICAL", "Out of stock (51 per board). Switch to FOJAN FRC0402F1002TS (C2906861, 14 M pcs). JLC basic part: check JLC's own SMT stock first (no extended-part fee)."),
 "C52923": ("CRITICAL", "Out of stock. Switch to CCTC TCC0402X5R105K500AT (C20539426, 1 uF 50 V X5R, 179 k pcs)."),
 "C22939": ("CRITICAL", "Out of stock, all UNI-ROYAL variants 0. Switch R62/R63 to FOJAN FRC0603F2R20TS (C2907006)."),
 "C25092": ("CRITICAL", "Out of stock. Switch R89 to FOJAN FRC0402F22R0TS (C2929994)."),
 "C190095": ("CRITICAL", "Out of stock. Switch R82 (VIN divider, 0.1 %) to Vishay TNPW040210K0BEED (C844452, 0.1 % 25 ppm)."),
 "C1002": ("CRITICAL", "GZ1608D601TF no longer listed on lcsc.com. Switch FB1 to Murata BLM18KG601SN1D (C85833, 600R@100MHz, 806 k pcs)."),
 "C86617": ("CRITICAL", "SWPA5040S470MT no longer listed on lcsc.com. Switch L2/L3 to APV ANR5040T470M (C7427133, 47 uH, Isat 1.3 A, 5x5 shielded); check DCR on the datasheet (peak current only 0.27 A)."),
 "C53134": ("CRITICAL", "BOM ERROR: C82 must be C0G (VIN-monitor ramp cap) but C53134 is X7R. Use KEMET C0805C473F3GECAUTO (C2233314, C0G 1 %, 25 V - ramp < 5 V)."),
 "C500738": ("HIGH", "LCSC marks LR2512D-3W-4mR DISCONTINUED (3,680 left). Switch RS1 to Milliohm HoJLR2512-3W-4mR-1% (C2904236, 21 k pcs, 3 W 2512)."),
 "C5718384": ("HIGH", "102 pcs, no drop-in at LCSC. Buy now. 2nd: PA3855.005NLT (520 pcs + 25 k from other suppliers) is the same EP13/21 uH part with 5 V windings - series-connect its two secondaries (needs a PCB change, plan for Rev F)."),
 "C49308184": ("HIGH", "120 pcs. Buy now. Alternates are 17-18 mm 1770-size parts with low stock: check pad drawings against L_Bourns_SRP1770TA before use."),
 "C249664": ("HIGH", "175 pcs, 3 per board (58 boards). 2nd: ChengTao RVE1J470M1010 (exact 47 uF 63 V 10x10.5, 350 pcs); 3rd: JIERR 100 uF 63 V 10x10.5 (24,990 pcs) - more input C, check F1 inrush."),
 "C141980": ("HIGH", "338 pcs, 3 per board (112 boards). 2nd UCC27211AQDDARQ1 (same DDA package, 755 pcs); 3rd UCC27211DR (SOIC-8 w/o thermal pad, 2,138 pcs, fits the footprint)."),
 "C224219": ("MEDIUM", "1,360 pcs, 4 per board. Recommend switching primary to MDD SS110 (C2482, 163 k pcs, JLC basic part)."),
 "C2940749": ("MEDIUM", "1,802 pcs, 6 per board (300 boards). CCTC 100 nF 100 V C0G 1206 (C20539418, 10 k pcs) is the 2nd source."),
 "C206980": ("MEDIUM", "765 pcs. Reomax R12.000.7 (C5383192, 7 A 63 V DC 1206, 100 A breaking) verified as 2nd source."),
 "C891118": ("MEDIUM", "1,606 pcs. HUASHUO HSS2N15 (150 V, 1.8 A, SOT-23) and Cmos CMN2N15M (150 V) as alternates - check Vgs(th) vs UC3843 10 V drive (fine)."),
 "C870479": ("MEDIUM", "Only one 191k 0.1 % 1206 in stock on LCSC (4,755 pcs, 1 per board); no LCSC alternate - fallback Panasonic ERA-8AEB1913V / Vishay TNPW1206191KBEEA from Digikey/Mouser."),
 "C5338808": ("MEDIUM", "4,777 pcs, but no drop-in: CA-IS3062W (C2890050) has no integrated isolated supply. Keep stock in mind for large builds."),
 "C6331172": ("LOW", "11,262 pcs; only alternate C2859409 (35 pcs)."),
 "C2933247": ("LOW", "423 k pcs; Yageo RC0603FR-0760R4L 2nd; no 3rd 60.4R 0603 in stock (62R 1 % FOJAN would also do for CAN split termination)."),
 "C81598": ("LOW", "2.9 M pcs (JLC basic)."),
}

wb = Workbook()
ws = wb.active
ws.title = "Sourcing"
hdr = ["#", "Designators", "Qty / board", "Description", "Primary MPN", "Manufacturer", "LCSC #", "LCSC stock today",
       "Boards covered", "LCSC price from (USD, high-qty tier)", "Risk", "Action / note",
       "2nd source MPN", "2nd source note", "2nd LCSC #", "2nd LCSC stock",
       "3rd source MPN", "3rd source note", "3rd LCSC #", "3rd LCSC stock", "Sources in stock (>= qty for 100 boards)"]
ws.append(hdr)
F = Font(name="Arial", size=9)
B = Font(name="Arial", size=9, bold=True)
thin = Side(style="thin", color="BBBBBB")
fills = {"CRITICAL": "F4B6B6", "HIGH": "F9D7A8", "MEDIUM": "FFF2B3", "LOW": "D9EFD3"}
rows = []
for i, (comment, des, fp, lcsc) in enumerate(bom, 1):
    qty = len(des.split(","))
    mpn, mfr, st, fr = STOCK.get(lcsc, ("?", "?", "?", ""))
    stock = 0 if st in ("NOTLISTED", "?", "") else int(st)
    alts = ALTS.get(lcsc, [("", ""), ("", "")])
    a = []
    for code, note in alts[:2]:
        if code:
            m2, f2, s2, _ = STOCK.get(code, ("?", "?", "?", ""))
            a.append((f"{m2} ({f2})", note, code, 0 if s2 in ("NOTLISTED", "?") else int(s2)))
        else:
            a.append(("", "none in stock on LCSC", "", None))
    boards = stock // qty
    if lcsc in FIX:
        risk, note = FIX[lcsc]
    else:
        n_ok = sum(1 for s in [stock] + [x[3] or 0 for x in a] if s >= 100 * qty)
        if boards < 100:
            risk = "HIGH"
        elif boards < 1000 or n_ok < 2:
            risk = "MEDIUM"
        else:
            risk = "LOW"
        note = {"LOW": "Stock fine.", "MEDIUM": "Primary stock covers the build; alternates have little LCSC stock - keep an eye on it.", "HIGH": "Low stock for the qty per board - buy now."}[risk]
    rows.append([i, des, qty, comment, mpn, mfr, lcsc, st if st == "NOTLISTED" else stock, None, float(fr) if fr else None,
                 risk, note] + [a[0][0], a[0][1], a[0][2], a[0][3], a[1][0], a[1][1], a[1][2], a[1][3]])
order = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3}
rows.sort(key=lambda r: (order[r[10]], r[0]))
for n, r in enumerate(rows, 2):
    ws.append(r)
    ws.cell(n, 9).value = f'=IF(ISNUMBER(H{n}),INT(H{n}/C{n}),0)'
    ws.cell(n, 21).value = (f'=IF(ISNUMBER(H{n}),IF(H{n}>=100*C{n},1,0),0)+IF(ISNUMBER(P{n}),IF(P{n}>=100*C{n},1,0),0)'
                            f'+IF(ISNUMBER(T{n}),IF(T{n}>=100*C{n},1,0),0)')
    for c in range(1, len(hdr) + 1):
        cell = ws.cell(n, c)
        cell.font = F
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        cell.alignment = Alignment(vertical="top", wrap_text=c in (2, 4, 12, 14, 18))
    ws.cell(n, 11).fill = PatternFill("solid", fgColor=fills[r[10]])
    ws.cell(n, 11).font = B
    for c in (8, 9, 16, 20):
        ws.cell(n, c).number_format = "#,##0"
    ws.cell(n, 10).number_format = "0.0000"
for c in range(1, len(hdr) + 1):
    h = ws.cell(1, c)
    h.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
    h.fill = PatternFill("solid", fgColor="2F4F6F")
    h.alignment = Alignment(wrap_text=True, vertical="center")
widths = [4, 22, 6, 18, 22, 16, 11, 10, 9, 9, 10, 60, 26, 30, 11, 10, 26, 30, 11, 10, 11]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{len(rows)+1}"
ws.row_dimensions[1].height = 42

# ---------------- Summary
s = wb.create_sheet("Summary", 0)
lines = [
 ["DCDC_LLC_4SBB_G474 Rev E - material shortage review and 2nd / 3rd sources"],
 ["LCSC stock read live from lcsc.com product pages on 2026-10-20 (about 09:00 Beijing time). Stock moves daily - re-check at order time."],
 ["JLC assembly uses JLC's own SMT library: a JLC 'basic' part can still be in stock there when LCSC retail shows 0, and each extra 'extended' part adds JLC's per-part setup fee."],
 [],
 ["Risk", "Lines", "Rule"],
 ["CRITICAL", '=COUNTIF(Sourcing!K:K,"CRITICAL")', "Out of stock / not listed / < 50 boards, or a BOM error - must change before ordering"],
 ["HIGH", '=COUNTIF(Sourcing!K:K,"HIGH")', "Covers < 100-300 boards, single or weak source, or discontinued - buy now / plan alternate"],
 ["MEDIUM", '=COUNTIF(Sourcing!K:K,"MEDIUM")', "Covers < 1,000 boards or only one source with stock for 100 boards"],
 ["LOW", '=COUNTIF(Sourcing!K:K,"LOW")', "Plenty of stock and at least two sources"],
 [],
 ["Must do before the Rev E order (BOM / JLC BOM changes - same footprints):"],
]
for l in lines:
    s.append(l)
acts = [r for r in rows if r[10] == "CRITICAL"]
for r in acts:
    s.append([r[6], r[1], r[11]])
s.append([])
s.append(["Buy now / watch:"])
for r in rows:
    if r[10] == "HIGH":
        s.append([r[6], r[1], r[11]])
s.append([])
s.append(["Hand-soldered / not on the JLC BOM: J1, J2 (KF2EDGR-5.08-2P), J3, J4, J5, J7, JP1 headers - commodity; T1 custom PQ2620 (see winding guide); U1 is listed on the Sourcing sheet because it is the main shortage item."])
s.column_dimensions["A"].width = 14
s.column_dimensions["B"].width = 34
s.column_dimensions["C"].width = 120
for row in s.iter_rows():
    for c in row:
        c.font = Font(name="Arial", size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
s["A1"].font = Font(name="Arial", size=13, bold=True)
for rr in range(6, 10):
    s.cell(rr, 1).fill = PatternFill("solid", fgColor=fills[s.cell(rr, 1).value])
    s.cell(rr, 1).font = Font(name="Arial", size=10, bold=True)
for rr in (5, 11):
    for c in range(1, 4):
        s.cell(rr, c).font = Font(name="Arial", size=10, bold=True)
out = sys.argv[1]
wb.save(out)
print(len(rows), "lines;", {k: sum(1 for r in rows if r[10] == k) for k in order})
